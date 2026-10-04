"""Export one board to CSV or Excel, in the column order of Mahak's existing workbook: the Zoho deal
fields, then each action (objective, action, why now, evidence, proof, SME), then ticked, rationale and
own action.

One row per action (deal fields repeat), or one row for a deal with no actions yet. Ticks, rationale and
own action are the current decision for the week the actions were drafted in; undecided, they are blank. Cells that a
spreadsheet would read as a formula are neutralised, because rationale and own action are free text.
"""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy.orm import Session

from api.domain.stages import stage_position
from api.models import Deal
from api.views import SOURCE_LABEL, visible_deals
from api.weekly import latest_actions, reviews_for

COLUMNS = ["Potentials Owner", "Potentials Name", "Stage", "City/State", "Potential Source", "Industry Type",
           "EP Involved", "EL Involved", "Objective", "Action", "Why now", "Evidence", "Proof", "SME",
           "Ticked", "Rationale", "Own action"]
OBJECTIVE_NAME = {"advance": "Advance", "unblock": "Unblock", "reframe": "Reframe", "nurture": "Nurture",
                  "re_engage": "Re-engage"}
FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def safe(value):
    """A cell value that no spreadsheet will run as a formula."""
    if isinstance(value, str) and value.startswith(FORMULA_STARTS):
        return "'" + value
    return value


def evidence_text(evidence: list[dict]) -> str:
    """Each fact the action rests on, as "Source: what (date)", one per line of a cell."""
    return "\n".join(f"{SOURCE_LABEL.get(e.get('source'), e.get('source'))}: {e.get('ref')}"
                     + (f" ({e['date']})" if e.get("date") else "") for e in evidence or [])


def proof_text(proof: dict | None) -> str | None:
    if not proof or not proof.get("name"):
        return None
    return proof["name"] + (f": {proof['why']}" if proof.get("why") else "")


def rows_for(session: Session, board: str) -> list[list]:
    deals = [d for d in session.scalars(visible_deals().where(Deal.board == board))]
    deals.sort(key=lambda d: (-(stage_position(d.stage) or 0), d.name.lower()))
    out = []
    for d in deals:
        week, recs = latest_actions(session, d.id)
        review, picks = reviews_for(session, d.id).get(week, (None, {})) if week else (None, {})
        head = [d.owner_name, d.name, d.stage, d.city_state, d.lead_source, d.industry,
                ", ".join(d.ep_involved or []), ", ".join(d.el_involved or [])]
        decision = [review.rationale if review else None, review.own_action if review else None]
        if not recs:
            out.append(head + [None] * 6 + [None] + decision)
            continue
        for r in recs:
            ticked = None if review is None else ("Yes" if picks.get(r.id) else "No")
            out.append(head + [OBJECTIVE_NAME.get(r.objective, r.objective), r.action, r.why_now,
                               evidence_text(r.evidence), proof_text(r.proof), r.sme, ticked] + decision)
    return [[safe(v) for v in row] for row in out]


def to_csv(rows: list[list]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    writer.writerows(rows)
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")  # a BOM so Excel reads the accents


def to_xlsx(rows: list[list], title: str) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")  # evidence holds one fact per line
    sheet.freeze_panes = "A2"
    for index, name in enumerate(COLUMNS, 1):
        width = 60 if name in ("Action", "Why now", "Evidence", "Rationale", "Own action") else 22
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
