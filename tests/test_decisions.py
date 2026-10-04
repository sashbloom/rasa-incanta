"""The decision layer: saving a decision, the weekly history, the summary and the board export.

Seeded runs are dated NOW (Thursday 24 Sep 2026, week of 21 Sep); later weeks are runs dated NOW + n weeks.
Only the top deal (Northwind) gets an NBA from a default run, and the fixture reply is a single action, so
tests that need several actions add them to that run directly."""
import csv
import io
import uuid
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from api.config import get_settings
from api.db import get_sessionmaker
from api.engine.run import run_week
from api.models import Deal, DealReview, Decision, Recommendation
from api.sources.zoho import fetch_deals
from api.views import summary_view
from tests.fakes import NOW, FakeClaude, FakeConnection, claude_response, load, settings, zoho_connect

WEEK_1 = "2026-09-21"
EXPORT_COLUMNS = ["Potentials Owner", "Potentials Name", "Stage", "City/State", "Potential Source", "Industry Type",
                  "EP Involved", "EL Involved", "Objective", "Action", "Why now", "Evidence", "Proof", "SME",
                  "Ticked", "Rationale", "Own action"]  # the order of Mahak's workbook
RATIONALE = "The CFO owns the phasing question, so the plan comes first."


def reply():
    return claude_response(load("claude/nba_reply.json"))


def run_at(when, rows=None, replies=4):
    conn = FakeConnection(deals=rows) if rows is not None else FakeConnection()
    with get_sessionmaker()() as s:
        return run_week(s, settings(), now=when, fetch=lambda st: fetch_deals(st, connect_fn=zoho_connect(conn)),
                        llm_client=FakeClaude(*[reply() for _ in range(replies)]))


def deal_id(name):
    with get_sessionmaker()() as s:
        return str(s.scalar(select(Deal.id).where(Deal.name.startswith(name))))


def add_actions(name, n=2):
    """Extra actions on a deal's latest run, so a decision has something to tick and leave."""
    with get_sessionmaker()() as s:
        base = s.scalars(select(Recommendation).join(Deal, Deal.id == Recommendation.deal_id)
                         .where(Deal.name.startswith(name)).order_by(Recommendation.created_at.desc())).first()
        for i in range(n):
            s.add(Recommendation(run_id=base.run_id, deal_id=base.deal_id, rank=base.rank + i + 1,
                                 objective=["nurture", "reframe"][i % 2], action=f"Extra action {i + 1}",
                                 why_now="Because.", evidence=base.evidence, gaps=[], model="test"))
        s.commit()


@pytest.fixture()
def seeded(client):
    run_at(NOW)
    add_actions("Northwind")
    return client


def detail(client, name="Northwind"):
    return client.get(f"/api/deals/{deal_id(name)}").json()


def decide(client, ticked, rationale=RATIONALE, own=None, name="Northwind"):
    d = detail(client, name)
    ids = [d["actions"][i]["id"] for i in ticked]
    return client.post(f"/api/deals/{d['id']}/decisions", json={"selected": ids, "rationale": rationale, "own_action": own})


def count(model):
    with get_sessionmaker()() as s:
        return s.scalar(select(func.count()).select_from(model))


# ---------------------------------------------------------------- saving

def test_a_decision_stores_the_ticks_the_rationale_and_the_own_action(seeded):
    response = decide(seeded, [0, 2], own="Call the CFO myself on Friday.")
    assert response.status_code == 201
    body = response.json()
    assert [a["selected"] for a in body["actions"]] == [True, False, True]
    assert body["decision"]["rationale"] == RATIONALE
    assert body["decision"]["own_action"] == "Call the CFO myself on Friday."
    assert body["decision"]["week_start"] == WEEK_1
    with get_sessionmaker()() as s:
        review = s.scalar(select(DealReview))
        assert review.week_start.isoformat() == WEEK_1 and review.user_id is None and review.status == "done"
        decisions = list(s.scalars(select(Decision).where(Decision.review_id == review.id)))
    assert len(decisions) == 3 and sum(d.selected for d in decisions) == 2  # unticked actions are recorded too


def test_nothing_ticked_is_a_valid_decision_with_a_rationale(seeded):
    response = decide(seeded, [], rationale="None of these fit; the deal is waiting on legal.")
    assert response.status_code == 201 and [a["selected"] for a in response.json()["actions"]] == [False] * 3


def test_a_rationale_is_required(seeded):
    for blank in ("", "   \n"):
        response = decide(seeded, [0], rationale=blank)
        assert response.status_code == 422 and "rationale" in response.json()["detail"]
    assert count(DealReview) == 0 and count(Decision) == 0


def test_only_the_deals_own_actions_can_be_ticked(seeded):
    d = detail(seeded)
    other = client_post(seeded, d["id"], [str(uuid.uuid4())])
    assert other.status_code == 422
    assert count(DealReview) == 0


def client_post(client, deal, ids):
    return client.post(f"/api/deals/{deal}/decisions", json={"selected": ids, "rationale": RATIONALE})


def test_editing_appends_a_new_version_and_the_old_one_stays(seeded):
    assert decide(seeded, [0], rationale="First thought.").status_code == 201
    edited = decide(seeded, [1, 2], rationale="Changed my mind.", own="Call the CFO.")
    assert edited.status_code == 201
    body = edited.json()
    assert body["decision"]["version"] == 2 and body["decision"]["rationale"] == "Changed my mind."
    assert [a["selected"] for a in body["actions"]] == [False, True, True]
    assert body["decision"]["own_action"] == "Call the CFO."
    earlier = body["decision"]["earlier"]
    assert [(e["version"], e["rationale"], len(e["selected_ids"])) for e in earlier] == [(1, "First thought.", 1)]
    with get_sessionmaker()() as s:
        reviews = list(s.scalars(select(DealReview).order_by(DealReview.version)))
        assert [(r.version, r.rationale) for r in reviews] == [(1, "First thought."), (2, "Changed my mind.")]
        assert len({r.week_start for r in reviews}) == 1  # same deal, same week
    assert count(Decision) == 6  # three actions, ticked or not, per version: nothing was overwritten


def test_the_latest_version_is_current_everywhere(seeded):
    decide(seeded, [0], rationale="First thought.")
    decide(seeded, [2], rationale="Second thought.")
    decide(seeded, [1], rationale="Third thought.")
    d = detail(seeded)
    assert d["decision"]["version"] == 3 and d["decision"]["rationale"] == "Third thought."
    assert [e["version"] for e in d["decision"]["earlier"]] == [2, 1]
    rows = list(csv.DictReader(io.StringIO(seeded.get("/api/export?board=pipeline&format=csv").content.decode("utf-8-sig"))))
    northwind = [r for r in rows if r["Potentials Name"].startswith("Northwind")]
    assert [r["Ticked"] for r in northwind] == ["No", "Yes", "No"] and {r["Rationale"] for r in northwind} == {"Third thought."}


def test_an_edit_still_needs_a_rationale(seeded):
    decide(seeded, [0])
    assert decide(seeded, [1], rationale=" ").status_code == 422
    assert count(DealReview) == 1


def test_two_edits_cannot_claim_the_same_version(seeded):
    from sqlalchemy.exc import IntegrityError
    decide(seeded, [0])
    with get_sessionmaker()() as s:
        first = s.scalar(select(DealReview))
        s.add(DealReview(week_start=first.week_start, deal_id=first.deal_id, user_id=None, version=1, rationale="x"))
        with pytest.raises(IntegrityError):
            s.commit()


def test_history_shows_the_current_version_of_each_week(client):
    run_at(NOW)
    add_actions("Northwind")
    decide(client, [0], rationale="First thought.")
    decide(client, [2], rationale="Second thought.")
    run_at(NOW + timedelta(weeks=1))
    old = detail(client)["history"][0]
    assert old["rationale"] == "Second thought." and [a["selected"] for a in old["actions"]] == [False, False, True]


def test_a_deal_with_no_actions_cannot_be_decided(seeded):
    response = client_post(seeded, deal_id("Blue Harbour"), [])
    assert response.status_code == 409


def test_an_unknown_deal_is_a_404(seeded):
    assert client_post(seeded, str(uuid.uuid4()), []).status_code == 404


def test_the_board_marks_a_decided_deal(seeded):
    def flags():
        boards = seeded.get("/api/week").json()["boards"]
        return {d["name"].split(" - ")[0]: d["decided"] for b in boards for st in b["stages"] for d in st["deals"]}

    assert flags() == {"Northwind Foods": False, "Blue Harbour Retail": False}
    decide(seeded, [0])
    assert flags() == {"Northwind Foods": True, "Blue Harbour Retail": False}


def test_a_rerun_leaves_a_decided_deals_actions_alone(seeded):
    decide(seeded, [0])
    before = count(Recommendation)
    run_at(NOW)  # same week
    assert count(Recommendation) == before  # nothing new for Northwind; Blue Harbour still has no NBA by default


# ---------------------------------------------------------------- history

def test_weekly_history_keeps_every_week_with_its_ticks_and_rationale(client):
    run_at(NOW)
    add_actions("Northwind")
    decide(client, [1])
    for week in range(1, 7):  # six more weeks of undecided actions
        run_at(NOW + timedelta(weeks=week))
    d = detail(client)
    assert d["actions_week"] == "2026-11-02"
    assert d["decision"] is None and all(a["selected"] is None for a in d["actions"])
    weeks = [h["week_start"] for h in d["history"]]
    assert weeks == sorted(weeks, reverse=True) and len(weeks) == 6 and weeks[-1] == WEEK_1  # nothing dropped
    oldest = d["history"][-1]
    assert oldest["decided"] and oldest["rationale"] == RATIONALE
    assert [a["selected"] for a in oldest["actions"]] == [False, True, False]
    newer = d["history"][0]
    assert newer["decided"] is False and all(a["selected"] is None for a in newer["actions"])  # undecided: unknown


def test_history_is_never_deleted_by_a_later_week(client):
    run_at(NOW)
    add_actions("Northwind")
    decide(client, [0])
    run_at(NOW + timedelta(weeks=1))
    assert count(DealReview) == 1 and count(Decision) == 3


# ---------------------------------------------------------------- summary

def moved_and_new_rows():
    rows = [dict(r) for r in load("zoho/deals.json")]
    rows[1]["stage"] = "Proposal Sent"  # Blue Harbour moves up
    rows.append(dict(rows[0], id=598723000011234999, deal_name="Contoso Steel - Supply chain",
                     _rel_account_name="Contoso Steel Ltd", stage="Qualified Prospect"))
    return rows


def test_first_week_has_nothing_to_compare_with(client):
    run_at(NOW)
    with get_sessionmaker()() as s:
        summary = summary_view(s, get_settings(), now=NOW)
    assert summary["comparison"] is False and summary["new"]["count"] == 0 and summary["moved"]["count"] == 0
    assert summary["pending"]["count"] == 1  # Northwind has an undecided action


def test_the_summary_counts_new_deals_stage_moves_and_pending_decisions(client):
    run_at(NOW)
    run_at(NOW + timedelta(weeks=1), rows=moved_and_new_rows())
    with get_sessionmaker()() as s:
        summary = summary_view(s, get_settings(), now=NOW + timedelta(weeks=1))
        week = client.get("/api/week")
    assert summary["comparison"] is True
    assert [d["name"] for d in summary["new"]["deals"]] == ["Contoso Steel - Supply chain"]
    assert [(d["name"], d["from"], d["to"]) for d in summary["moved"]["deals"]] == [
        ("Blue Harbour Retail - Cost audit", "Qualified Prospect", "Proposal Sent")]
    assert summary["pending"]["count"] == 1 and summary["decided"] == 0
    assert week.status_code == 200


def test_deciding_moves_a_deal_out_of_pending(client):
    run_at(NOW)
    decide(client, [0])
    with get_sessionmaker()() as s:
        summary = summary_view(s, get_settings(), now=NOW)
    assert summary["pending"]["count"] == 0 and summary["decided"] == 1


def test_the_summary_endpoint_answers(client):
    run_at(NOW)
    body = client.get("/api/summary").json()
    assert set(body) >= {"week_start", "comparison", "new", "moved", "pending", "decided"}


def test_the_board_flags_match_the_summary_so_each_count_links_to_its_list(client, monkeypatch):
    run_at(NOW)
    run_at(NOW + timedelta(weeks=1), rows=moved_and_new_rows())
    with get_sessionmaker()() as s:
        from api.views import week_view
        week = week_view(s, get_settings(), now=NOW + timedelta(weeks=1))
        summary = summary_view(s, get_settings(), now=NOW + timedelta(weeks=1))
    deals = [d for b in week["boards"] for st in b["stages"] for d in st["deals"]]
    assert {d["id"] for d in deals if d["new"]} == {d["id"] for d in summary["new"]["deals"]}
    assert {d["id"] for d in deals if d["moved"]} == {d["id"] for d in summary["moved"]["deals"]}
    assert {d["id"] for d in deals if d["has_actions"] and not d["decided"]} == {d["id"] for d in summary["pending"]["deals"]}


# ---------------------------------------------------------------- export

def test_the_csv_has_the_deal_fields_actions_ticks_and_rationale(seeded):
    decide(seeded, [0], own="Call the CFO myself.")
    response = seeded.get("/api/export?board=pipeline&format=csv")
    assert response.status_code == 200 and response.headers["content-type"].startswith("text/csv")
    assert 'filename="rasa-incanta-pipeline-' in response.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    northwind = [r for r in rows if r["Potentials Name"].startswith("Northwind")]
    assert len(northwind) == 3 and [r["Ticked"] for r in northwind] == ["Yes", "No", "No"]
    assert all(r["Rationale"] == RATIONALE for r in northwind)
    assert northwind[0]["Own action"] == "Call the CFO myself." and northwind[0]["Stage"] == "Proposal Sent"
    blue = next(r for r in rows if r["Potentials Name"].startswith("Blue Harbour"))
    assert blue["Action"] == "" and blue["Ticked"] == ""  # no actions yet: still listed


def test_an_undecided_deals_ticks_are_blank(seeded):
    rows = list(csv.DictReader(io.StringIO(seeded.get("/api/export?board=pipeline&format=csv").content.decode("utf-8-sig"))))
    assert {r["Ticked"] for r in rows if r["Potentials Name"].startswith("Northwind")} == {""}


def test_the_excel_export_opens_and_has_a_header_row(seeded):
    decide(seeded, [1])
    response = seeded.get("/api/export?board=pipeline")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [c.value for c in sheet[1]]== EXPORT_COLUMNS and sheet.max_row == 5


def test_free_text_that_looks_like_a_formula_is_neutralised(seeded):
    decide(seeded, [0], rationale='=HYPERLINK("http://evil.example","x")', own="+cmd|' /C calc'!A0")
    text = seeded.get("/api/export?board=pipeline&format=csv").content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'+cmd" in text
    sheet = load_workbook(io.BytesIO(seeded.get("/api/export?board=pipeline").content)).active
    assert not any(isinstance(c.value, str) and c.value.startswith(("=", "+")) for row in sheet.iter_rows() for c in row)


def test_the_export_validates_its_board_and_format(seeded):
    assert seeded.get("/api/export?board=nope").status_code == 422
    assert seeded.get("/api/export?board=pipeline&format=pdf").status_code == 422
    assert seeded.get("/api/export").status_code == 422
    empty = seeded.get("/api/export?board=prospect&format=csv")
    assert empty.status_code == 200 and empty.content.decode("utf-8-sig").startswith(",".join(EXPORT_COLUMNS))


def test_the_summary_page_is_served_at_both_roots(client):
    for path in ("/summary", "/reports/rasa-incanta/summary"):
        response = client.get(path)
        assert response.status_code == 200 and '<div id="root">' in response.text


def test_the_export_follows_mahaks_column_order_and_fills_the_deal_fields(seeded):
    text = seeded.get("/api/export?board=pipeline&format=csv").content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    assert reader.fieldnames == EXPORT_COLUMNS
    first = next(r for r in reader if r["Potentials Name"].startswith("Northwind"))
    assert first["Potentials Owner"] == "Rao" and first["Stage"] == "Proposal Sent"
    assert first["City/State"] == "Pune, Maharashtra" and first["Potential Source"] == "Referral"
    assert first["Industry Type"] == "Food processing" and first["EP Involved"] == "A. Mehta"
    assert first["Objective"] == "Advance" and first["Evidence"].startswith("Zoho: ")
